"""
Excel Export Service for Camera-Trap Ingestion & Wildlife Analytics
Generates multi-sheet formatted Excel workbooks (.xlsx) with:
1. Executive Batch Overview (timestamps, total images, animal counts, reduction rate)
2. Species Breakdown & Frequency Counts ("how many are there")
3. Granular Image-by-Image Audit Log with individual timestamps, confidence, and sensor diagnostics
"""

import io
from datetime import datetime
from typing import Dict, Any, List, Optional
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from src.backend.services.alert_service import alert_service

class ExcelExportService:
    """
    Handles generation of styled Microsoft Excel (.xlsx) survey reports.
    """

    @staticmethod
    def generate_batch_excel_report(
        batch_data: Dict[str, Any],
        location: str = "loc_01",
        export_time: Optional[str] = None
    ) -> bytes:
        """
        Builds a multi-tab formatted Excel report for batch uploaded camera-trap frames.
        """
        if not export_time:
            export_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        items = batch_data.get("items", [])
        total_processed = len(items)
        empty_count = sum(1 for item in items if item.get("is_empty", False))
        animal_count = total_processed - empty_count

        # 1. Species aggregation ("how many are there")
        species_stats: Dict[str, Dict[str, Any]] = {}
        for item in items:
            if not item.get("is_empty", False):
                sp = item.get("common_name") or item.get("species", "Unknown")
                sp_key = item.get("species", sp)
                conf = float(item.get("confidence", 0.0))
                if sp not in species_stats:
                    species_stats[sp] = {"count": 0, "conf_sum": 0.0, "species_key": sp_key}
                species_stats[sp]["count"] += 1
                species_stats[sp]["conf_sum"] += conf

        rare_count = sum(1 for item in items if item.get("is_rare") or (item.get("alert") and item["alert"].get("is_rare")))
        dangerous_count = sum(1 for item in items if item.get("is_dangerous") or (item.get("alert") and item["alert"].get("is_dangerous")))
        duplicate_count = sum(1 for item in items if item.get("is_duplicate"))
        cluster_ids = set(item.get("cluster_id") for item in items if item.get("cluster_id"))
        unique_encounters = len(cluster_ids) if cluster_ids else (total_processed - duplicate_count)
        workload_reduction_str = f"{(duplicate_count / total_processed * 100):.1f}%" if total_processed else "0.0%"

        # Summary Tab
        summary_rows = [
            {"Metric": "Batch Export Timestamp", "Value": export_time},
            {"Metric": "Camera Station / Location", "Value": location},
            {"Metric": "Total Images Uploaded", "Value": total_processed},
            {"Metric": "Animals Detected (Count)", "Value": animal_count},
            {"Metric": "Empty Frames Filtered (Count)", "Value": empty_count},
            {"Metric": "Empty Reduction Rate (%)", "Value": f"{(empty_count / total_processed * 100):.1f}%" if total_processed else "0.0%"},
            {"Metric": "Distinct Species Count", "Value": len(species_stats)},
            {"Metric": "🐾 Unique Wildlife Encounters", "Value": unique_encounters},
            {"Metric": "🔄 Duplicate Images Grouped", "Value": duplicate_count},
            {"Metric": "⚡ Review Workload Saved (%)", "Value": workload_reduction_str},
            {"Metric": "🚨 Rare Animal Alerts (Count)", "Value": rare_count},
            {"Metric": "⚠️ Dangerous Wildlife Alerts (Count)", "Value": dangerous_count},
            {"Metric": "Model Backbone", "Value": "MobileNetV2 ONNX + CLAHE Bilateral Preprocessing"},
            {"Metric": "Temporal Burst Pooling", "Value": "Active (Rolling Max Window = 3)"}
        ]
        df_summary = pd.DataFrame(summary_rows)

        # Species Counts Tab ("how many are there")
        species_rows = []
        for sp_name, stats in sorted(species_stats.items(), key=lambda x: x[1]["count"], reverse=True):
            count = stats["count"]
            pct = (count / animal_count * 100) if animal_count else 0
            avg_conf = (stats["conf_sum"] / count * 100) if count else 0
            
            # Lookup alert status for this species
            alert_stat = alert_service.evaluate_animal_alert(stats.get("species_key", sp_name), common_name=sp_name)
            alert_badge = alert_stat.get("badge", "Normal Wildlife") if alert_stat.get("has_alert") else "Normal Wildlife"

            species_rows.append({
                "Species / Common Name": sp_name,
                "Count (How Many)": count,
                "Percentage of Wildlife": f"{pct:.1f}%",
                "Average Confidence (%)": f"{avg_conf:.1f}%",
                "Rarity & Threat Status": alert_badge
            })
        if not species_rows:
            species_rows = [{"Species / Common Name": "No Wildlife Detected", "Count (How Many)": 0, "Percentage of Wildlife": "0.0%", "Average Confidence (%)": "N/A", "Rarity & Threat Status": "None"}]
        df_species = pd.DataFrame(species_rows)

        # Image Records & Timestamps Tab
        detail_rows = []
        for idx, item in enumerate(items, 1):
            is_animal = not item.get("is_empty", False)
            item_time = item.get("datetime") or item.get("timestamp") or export_time
            bbox = item.get("bounding_box", {})
            if isinstance(bbox, dict):
                bbox_str = f"[{bbox.get('ymin', 0.0):.2f}, {bbox.get('xmin', 0.0):.2f}, {bbox.get('ymax', 0.0):.2f}, {bbox.get('xmax', 0.0):.2f}]"
            elif isinstance(bbox, list):
                bbox_str = str(bbox)
            else:
                bbox_str = "None"

            alert_obj = item.get("alert")
            if not alert_obj and is_animal:
                alert_obj = alert_service.evaluate_animal_alert(
                    species_key=item.get("species", "empty"),
                    common_name=item.get("common_name", ""),
                    confidence=float(item.get("confidence", 0.0))
                )
            alert_status = alert_obj.get("badge", "Normal") if alert_obj and alert_obj.get("has_alert") else "Normal"
            threat_level = alert_obj.get("threat_level", "Harmless") if alert_obj else "Harmless"

            # Frame duplicate role
            if item.get("is_anchor"):
                frame_role = "⭐ Anchor (Representative Frame)"
            elif item.get("is_duplicate"):
                frame_role = f"🔄 Duplicate (of {item.get('duplicate_of', 'Anchor')})"
            else:
                frame_role = "Single Encounter"

            sim_anchor = item.get("similarity_to_anchor")
            sim_str = f"{round(sim_anchor * 100, 1)}%" if sim_anchor is not None else ("100.0%" if item.get("is_anchor") else "N/A")

            detail_rows.append({
                "Record #": idx,
                "Timestamp": item_time,
                "File Name": item.get("file_name", "Unknown"),
                "Sighting Cluster": item.get("cluster_id", "Single Encounter"),
                "Frame Role": frame_role,
                "Similarity to Anchor (%)": sim_str,
                "Detection Status": "Animal Detected" if is_animal else "Empty Frame Filtered",
                "Species Code": item.get("species", "empty") if is_animal else "empty",
                "Common Name": item.get("common_name", "Empty Frame") if is_animal else "Animal Not Detected",
                "Scientific Name": item.get("scientific_name", "Identified Wildlife" if is_animal else "Non-Wildlife / Empty"),
                "Confidence (%)": round(float(item.get("confidence", 0.0)) * 100, 1),
                "Alert Status": alert_status,
                "Threat Level": threat_level,
                "Camera Station": item.get("location", location),
                "Night Vision (IR)": "Yes" if item.get("is_night") else "No",
                "Bounding Box": bbox_str if is_animal else "[0, 0, 0, 0]",
                "Uncertainty Flags": ", ".join(item.get("uncertainty_flags", [])) if item.get("uncertainty_flags") else "Normal",
                "Diagnostic Reason": item.get("empty_reason", "") or ("Animal Sighting Verified" if is_animal else "Ambient Motion Trigger")
            })
        df_details = pd.DataFrame(detail_rows)

        # Stream into memory with OpenPyXL
        output = io.BytesIO()
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df_summary.to_excel(writer, sheet_name="Batch Overview", index=False)
            df_species.to_excel(writer, sheet_name="Species Counts", index=False)
            df_details.to_excel(writer, sheet_name="Image Records & Timestamps", index=False)

            wb = writer.book
            header_fill = PatternFill(start_color="107C41", end_color="107C41", fill_type="solid")
            header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
            thin_border = Border(
                left=Side(style='thin', color='E0E0E0'),
                right=Side(style='thin', color='E0E0E0'),
                top=Side(style='thin', color='E0E0E0'),
                bottom=Side(style='thin', color='E0E0E0')
            )

            for sheetname in wb.sheetnames:
                ws = wb[sheetname]
                # Header row styling
                for cell in ws[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = Alignment(horizontal="center", vertical="center")

                # Column widths and cell borders
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        cell.border = thin_border
                        val_str = str(cell.value or "")
                        if len(val_str) > max_len:
                            max_len = len(val_str)
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

        output.seek(0)
        return output.getvalue()

excel_service = ExcelExportService()
