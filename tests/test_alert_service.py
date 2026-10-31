"""
Unit & Integration Tests for Wildlife Rarity & Threat Alert Service
Verifies:
1. Rare & dangerous animal classification (Mountain Lion, Wolf, Bear, Tiger, Elephant).
2. Rare non-dangerous animal classification (Gray Fox).
3. Benign common wildlife classification (Mule Deer, Rabbit, Bird, Empty).
4. Batch alert compilation and advisory generation.
5. REST API endpoint responses with alert metadata and GET /api/detection/alerts.
6. Excel export inclusion of alert metrics and threat status.
"""

import pytest
from io import BytesIO
from PIL import Image
from fastapi.testclient import TestClient

from src.backend.services.alert_service import alert_service
from src.backend.services.excel_service import excel_service
from src.backend.main import app

client = TestClient(app)

def test_mountain_lion_rare_and_dangerous_alert():
    """Mountain lion must trigger a CRITICAL Rare & Dangerous alert."""
    alert = alert_service.evaluate_animal_alert(
        species_key="mountain_lion",
        common_name="Mountain Lion",
        confidence=0.92,
        location="loc_01"
    )
    assert alert["has_alert"] is True
    assert alert["is_rare"] is True
    assert alert["is_dangerous"] is True
    assert alert["alert_level"] == "CRITICAL"
    assert alert["alert_type"] == "RARE_AND_DANGEROUS"
    assert "🚨" in alert["badge"]
    assert "standoff distance" in alert["safety_protocol"].lower()

def test_wolves_rare_and_dangerous_alert():
    """Alaskan tundra wolf and gray wolf must trigger CRITICAL Rare & Dangerous alert."""
    alert_tundra = alert_service.evaluate_animal_alert(
        species_key="alaskan_tundra_wolf",
        common_name="Alaskan Tundra Wolf",
        confidence=0.89,
        location="loc_02"
    )
    assert alert_tundra["has_alert"] is True
    assert alert_tundra["is_rare"] is True
    assert alert_tundra["is_dangerous"] is True
    assert alert_tundra["alert_level"] == "CRITICAL"

    alert_wolf = alert_service.evaluate_animal_alert(
        species_key="wolf",
        common_name="Wolf",
        confidence=0.88,
        location="loc_04"
    )
    assert alert_wolf["has_alert"] is True
    assert alert_wolf["is_dangerous"] is True

def test_bear_tiger_elephant_alerts():
    """Bears, tigers, and elephants must trigger CRITICAL alerts."""
    for sp in ["bear", "tiger", "elephant"]:
        alert = alert_service.evaluate_animal_alert(sp, sp.title(), 0.90)
        assert alert["has_alert"] is True
        assert alert["is_rare"] is True
        assert alert["is_dangerous"] is True
        assert alert["alert_level"] == "CRITICAL"

def test_gray_fox_rare_sighting_alert():
    """Gray Fox is rare but non-dangerous to humans."""
    alert = alert_service.evaluate_animal_alert(
        species_key="gray_fox",
        common_name="Gray Fox",
        confidence=0.85
    )
    assert alert["has_alert"] is True
    assert alert["is_rare"] is True
    assert alert["is_dangerous"] is False
    assert alert["alert_level"] == "WARNING"
    assert alert["alert_type"] == "RARE_SPECIES"

def test_benign_wildlife_no_critical_alert():
    """Mule deer, cottontail rabbit, bird, and empty frames do not trigger critical alerts."""
    for sp in ["mule_deer", "cottontail_rabbit", "bird", "rodent", "empty"]:
        alert = alert_service.evaluate_animal_alert(sp, sp.title(), 0.95)
        assert alert["is_dangerous"] is False
        assert alert["is_rare"] is False
        assert alert["alert_level"] == "NORMAL"

def test_batch_alerts_summary():
    """Summarizes batch detections and reports aggregate alert metrics."""
    batch_items = [
        {"species": "mountain_lion", "common_name": "Mountain Lion", "confidence": 0.92, "is_empty": False},
        {"species": "mule_deer", "common_name": "Mule Deer", "confidence": 0.95, "is_empty": False},
        {"species": "alaskan_tundra_wolf", "common_name": "Alaskan Tundra Wolf", "confidence": 0.88, "is_empty": False},
        {"species": "empty", "common_name": "Empty Frame", "confidence": 0.99, "is_empty": True}
    ]

    summary = alert_service.summarize_batch_alerts(batch_items)
    assert summary["has_alerts"] is True
    assert summary["rare_alerts_count"] == 2
    assert summary["dangerous_alerts_count"] == 2
    assert summary["critical_alerts_count"] == 2
    assert "Mountain Lion" in summary["alert_species"]
    assert "Alaskan Tundra Wolf" in summary["alert_species"]
    assert "🚨 CRITICAL WILDLIFE ALERT" in summary["headline"]

def test_api_predict_image_alert_metadata():
    """Verifies that POST /api/detection/predict-image returns alert fields."""
    img = Image.new("RGB", (256, 192), color=(40, 40, 40))
    buf = BytesIO()
    img.save(buf, format="JPEG")
    buf.seek(0)

    res = client.post(
        "/api/detection/predict-image",
        files={"file": ("test_alert.jpg", buf, "image/jpeg")},
        data={"location": "loc_01"}
    )
    assert res.status_code == 200
    data = res.json()
    assert "is_rare" in data
    assert "is_dangerous" in data
    assert "alert" in data
    assert isinstance(data["alert"], dict)
    assert "has_alert" in data["alert"]

def test_api_predict_batch_alert_metadata():
    """Verifies that POST /api/detection/predict-batch returns batch alert summary."""
    img = Image.new("RGB", (128, 128), color=(40, 40, 40))
    buf1 = BytesIO()
    img.save(buf1, format="JPEG")
    buf1.seek(0)
    buf2 = BytesIO()
    img.save(buf2, format="JPEG")
    buf2.seek(0)

    res = client.post(
        "/api/detection/predict-batch",
        files=[
            ("files", ("batch1.jpg", buf1, "image/jpeg")),
            ("files", ("batch2.jpg", buf2, "image/jpeg"))
        ],
        data={"location": "loc_01"}
    )
    assert res.status_code == 200
    data = res.json()
    assert "has_alerts" in data
    assert "rare_alerts_count" in data
    assert "dangerous_alerts_count" in data
    assert "alert_summary" in data
    assert len(data["items"]) == 2
    for it in data["items"]:
        assert "is_rare" in it
        assert "is_dangerous" in it
        assert "alert" in it

def test_api_get_alerts_endpoint():
    """Verifies GET /api/detection/alerts endpoint."""
    res = client.get("/api/detection/alerts")
    assert res.status_code == 200
    data = res.json()
    assert "total_alerts" in data
    assert "alerts" in data
    assert isinstance(data["alerts"], list)

def test_excel_export_includes_alert_metrics():
    """Verifies that Excel export includes Rare Animal and Dangerous Wildlife rows and columns."""
    sample_items = [
        {
            "file_name": "mt_lion.jpg",
            "species": "mountain_lion",
            "common_name": "Mountain Lion",
            "confidence": 0.94,
            "is_empty": False,
            "is_rare": True,
            "is_dangerous": True,
            "alert": alert_service.evaluate_animal_alert("mountain_lion", "Mountain Lion", 0.94)
        },
        {
            "file_name": "deer.jpg",
            "species": "mule_deer",
            "common_name": "Mule Deer",
            "confidence": 0.96,
            "is_empty": False,
            "is_rare": False,
            "is_dangerous": False
        }
    ]

    excel_bytes = excel_service.generate_batch_excel_report(
        batch_data={"items": sample_items},
        location="loc_01"
    )
    assert len(excel_bytes) > 2000

    # Read back with openpyxl to verify content
    import openpyxl
    wb = openpyxl.load_workbook(BytesIO(excel_bytes))
    overview_sheet = wb["Batch Overview"]
    metrics_found = [row[0] for row in overview_sheet.iter_rows(values_only=True)]
    assert any("Rare Animal Alerts" in str(m) for m in metrics_found)
    assert any("Dangerous Wildlife Alerts" in str(m) for m in metrics_found)

    records_sheet = wb["Image Records & Timestamps"]
    headers = [cell for cell in next(records_sheet.iter_rows(values_only=True))]
    assert "Alert Status" in headers
    assert "Threat Level" in headers
